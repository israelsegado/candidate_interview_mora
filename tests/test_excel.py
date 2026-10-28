"""Checks that results.xlsx exists and contains the same as results.csv."""
from openpyxl import load_workbook

from conftest import COLUMNS


def read_xlsx(path):
    sheet = load_workbook(path, read_only=True).worksheets[0]
    rows = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c.strip() for c in r)]
    header, data = rows[0], rows[1:]
    return [dict(zip(header, r)) for r in data]


def test_excel_exists(results_xlsx):
    assert results_xlsx.exists(), f"{results_xlsx.name} not found next to the CSV."


def test_excel_matches_csv(results_xlsx, rows):
    excel = read_xlsx(results_xlsx)
    assert len(excel) == len(rows), f"The Excel file has {len(excel)} rows and the CSV {len(rows)}."
    from_csv = {r["file"]: [r[c] for c in COLUMNS] for r in rows}
    errors = []
    for r in excel:
        values = [r.get(c, "") for c in COLUMNS]
        if from_csv.get(r.get("file")) != values:
            errors.append(r.get("file"))
    assert not errors, f"Excel rows that don't match the CSV: {errors}"
